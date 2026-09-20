"""Fit the supplied Lab 02 profiles with explicit width and baseline conventions.

The archived notebook uses different, unavailable TXT inputs. This script analyses
the four supplied XLSX files and records their hashes, residuals and limitations.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import platform
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
import scipy
from scipy.optimize import curve_fit

PROJECT = Path(__file__).resolve().parents[1]


def gaussian(x, amplitude, center, radius, offset):
    """Intensity model; radius is the 1/e² radius above the fitted baseline."""
    return amplitude * np.exp(-2 * ((np.asarray(x) - center) / radius) ** 2) + offset


def validate_profile(x, y):
    x, y = np.asarray(x, dtype=float), np.asarray(y, dtype=float)
    if x.ndim != 1 or y.shape != x.shape or x.size < 8:
        raise ValueError("A profile needs at least eight paired, one-dimensional samples.")
    if not np.isfinite(x).all() or not np.isfinite(y).all():
        raise ValueError("Profile values must all be finite.")
    if not np.all(np.diff(x) > 0):
        raise ValueError("Profile positions must be unique and strictly increasing.")
    if np.ptp(y) <= np.finfo(float).eps * max(1.0, float(np.max(np.abs(y)))):
        raise ValueError("Cannot fit a constant profile.")
    return x, y


def load_profile(path):
    """Read a headerless, single-column XLSX profile without dropping row one.

    Each row is one sample. Formulas, blanks, strings, booleans, additional sheets
    and additional columns are rejected instead of being silently reinterpreted.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Profile not found: {path}")
    if path.suffix.lower() != ".xlsx":
        raise ValueError("This importer expects a .xlsx profile.")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Expected exactly one worksheet.")
        sheet = workbook.worksheets[0]
        if sheet.max_column != 1:
            raise ValueError("Expected exactly one intensity column.")
        values = []
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            value = row[0]
            if isinstance(value, bool) or not isinstance(value, (int, float)):
                raise ValueError(f"Row {index} must contain a numeric intensity, not {value!r}.")
            values.append(value)
    finally:
        workbook.close()
    return validate_profile(np.arange(len(values), dtype=float), values)


def fit_profile(x, y, *, baseline="free", saturation_level=255.0):
    """Least squares with positive amplitude/radius and center inside the data.

    Baseline is either fitted freely (the notebook's model) or fixed to zero as
    a sensitivity check. Covariance errors assume independent residuals and do
    not include calibration, saturation, model mismatch or systematic errors.
    """
    x, y = validate_profile(x, y)
    if baseline not in {"free", "zero"}:
        raise ValueError("baseline must be 'free' or 'zero'.")
    if saturation_level is not None and (
        not np.isfinite(saturation_level) or saturation_level <= 0
    ):
        raise ValueError("saturation_level must be positive and finite, or None.")
    origin, span = float(x[0]), float(np.ptp(x))
    u = (x - origin) / span
    # Exclude the noise floor from the initial moment estimate. Otherwise a
    # narrow peak on a long noisy baseline can start with an enormous width.
    contrast = y - np.percentile(y, 20)
    weights = np.where(contrast > 0.1 * contrast.max(), contrast, 0.0)
    center0 = float(np.average(u, weights=weights))
    radius0 = max(2 * np.sqrt(np.average((u - center0) ** 2, weights=weights)), 1e-3)
    initial = [float(np.ptp(y)), center0, float(radius0)]
    lower, upper = [0, 0, 1e-8], [np.inf, 1, np.inf]
    if baseline == "free":
        model = gaussian
        initial.append(float(y.min()))
        lower.append(-np.inf)
        upper.append(np.inf)
    else:
        def model(t, a, c, w):
            return gaussian(t, a, c, w, 0)
    fitted, covariance = curve_fit(
        model, u, y, p0=initial, bounds=(lower, upper),
        x_scale="jac", max_nfev=30000,
    )
    amplitude, center_u, radius_u = fitted[:3]
    offset = float(fitted[3]) if baseline == "free" else 0.0
    center, radius = origin + center_u * span, radius_u * span
    predicted = gaussian(x, amplitude, center, radius, offset)
    residuals = y - predicted
    diagonal = np.diag(covariance)
    radius_stderr = float(np.sqrt(diagonal[2]) * span) if (
        np.isfinite(diagonal[2]) and diagonal[2] >= 0
    ) else None
    covered = bool(center - radius >= x[0] and center + radius <= x[-1])
    ceiling_count = int(np.count_nonzero(y >= saturation_level)) if saturation_level is not None else 0
    warnings = []
    if ceiling_count:
        warnings.append(f"{ceiling_count} samples reach the assumed {saturation_level:g}-count ceiling; possible clipping.")
    if not covered:
        warnings.append("The fitted 1/e² radius extends beyond the measured interval.")
    if offset < 0:
        warnings.append("The fitted baseline is negative; inspect baseline sensitivity and residuals.")
    if radius_stderr is None:
        warnings.append("A finite radius covariance error could not be estimated.")
    return {
        "baseline_mode": baseline,
        "samples": int(x.size),
        "amplitude_counts": float(amplitude),
        "center_samples": float(center),
        "radius_samples": float(radius),
        "diameter_samples": float(2 * radius),
        "fwhm_samples": float(np.sqrt(2 * np.log(2)) * radius),
        "radius_stderr_samples": radius_stderr,
        "offset_counts": offset,
        "rss_counts_squared": float(residuals @ residuals),
        "rmse_counts": float(np.sqrt(np.mean(residuals ** 2))),
        "full_radius_covered": covered,
        "samples_at_ceiling": ceiling_count,
        "warnings": warnings,
    }


def radius_ratio(before, after):
    """Ratio and conditional error, assuming independent fits and equal scales."""
    wb, wa = before["radius_samples"], after["radius_samples"]
    if not np.isfinite([wb, wa]).all() or min(wb, wa) <= 0:
        raise ValueError("Both radii must be positive and finite.")
    eb, ea = before["radius_stderr_samples"], after["radius_stderr_samples"]
    ratio = wa / wb
    error = None if eb is None or ea is None else ratio * np.hypot(eb / wb, ea / wa)
    return {"ratio": float(ratio), "conditional_stderr": None if error is None else float(error)}


def prediction(x, fit):
    return gaussian(x, fit["amplitude_counts"], fit["center_samples"],
                    fit["radius_samples"], fit["offset_counts"])


def save_figures(datasets, output):
    plt.rcParams.update({"font.size": 10, "axes.spines.top": False, "axes.spines.right": False})
    figure = plt.figure(figsize=(12, 9), layout="constrained")
    grid = figure.add_gridspec(4, 2, height_ratios=[2.3, 1, 2.3, 1])
    for index, (entry, x, y, free, zero) in enumerate(datasets):
        row, column = 2 * (index // 2), index % 2
        top = figure.add_subplot(grid[row, column])
        bottom = figure.add_subplot(grid[row + 1, column], sharex=top)
        top.plot(x, y, color="#64748b", linewidth=1, label="Measured profile")
        top.plot(x, prediction(x, free), color="#075985", linewidth=2, label="Gaussian + fitted baseline")
        top.set_title(f"{entry['label']}\nw = {free['radius_samples']:.2f} samples", loc="left")
        top.set_ylabel("Intensity (counts)")
        top.legend(fontsize=8, frameon=False)
        bottom.plot(x, y - prediction(x, free), color="#b45309", linewidth=0.8)
        bottom.axhline(0, color="#94a3b8", linewidth=0.8)
        bottom.set_ylabel("Residual")
        bottom.set_xlabel("Profile position (sample index)")
        # Zoom narrow lens profiles for readability; all samples enter the fit.
        if free["radius_samples"] < np.ptp(x) / 15:
            half = max(3 * free["radius_samples"], 30)
            top.set_xlim(max(x[0], free["center_samples"] - half),
                         min(x[-1], free["center_samples"] + half))
        for ax in (top, bottom):
            ax.grid(alpha=0.18)
    figure.suptitle("Lab 02 · Gaussian fits to the four supplied Excel profiles", fontsize=15)
    figure.savefig(output / "profile-fits.png", dpi=160)
    plt.close(figure)

    expander = [d for d in datasets if d[0]["id"] in {"before", "after"}]
    figure, axes = plt.subplots(1, len(expander), figsize=(12, 4.5), layout="constrained", squeeze=False)
    for ax, (entry, x, y, free, zero) in zip(axes.flat, expander):
        ax.plot(x, y, color="#94a3b8", linewidth=1, label="Measured profile")
        ax.plot(x, prediction(x, free), color="#075985", linewidth=2,
                label=f"Free baseline: w = {free['radius_samples']:.1f}")
        ax.plot(x, prediction(x, zero), color="#c2410c", linewidth=2, linestyle="--",
                label=f"Zero baseline: w = {zero['radius_samples']:.1f}")
        ax.set_title(entry["label"], loc="left")
        ax.set_xlabel("Profile position (sample index)")
        ax.set_ylabel("Intensity (counts)")
        ax.legend(fontsize=8, frameon=False)
        ax.grid(alpha=0.18)
    figure.suptitle("Baseline sensitivity · neither fit establishes the optical magnification", fontsize=13)
    figure.savefig(output / "expander-baseline-sensitivity.png", dpi=160)
    plt.close(figure)


def run_analysis(config_path, output_dir, *, make_figures=True):
    config_path, output = Path(config_path).resolve(), Path(output_dir).resolve()
    config = json.loads(config_path.read_text(encoding="utf-8"))
    # Paths belong to the config, so copied/custom configurations remain portable.
    entries = config["profiles"]
    identifiers = [entry["id"] for entry in entries]
    if len(entries) != 4 or len(set(identifiers)) != 4 or not {"before", "after"}.issubset(identifiers):
        raise ValueError("Configure four unique profiles, including 'before' and 'after'.")
    datasets, records = [], []
    for entry in entries:
        path = (config_path.parent / entry["file"]).resolve()
        x, y = load_profile(path)
        free = fit_profile(x, y, baseline="free", saturation_level=config["saturation_level_counts"])
        zero = fit_profile(x, y, baseline="zero", saturation_level=config["saturation_level_counts"])
        records.append({
            "id": entry["id"], "label": entry["label"], "file": entry["file"],
            "input_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "free_baseline": free, "zero_baseline_sensitivity": zero,
        })
        datasets.append((entry, x, y, free, zero))
    by_id = {record["id"]: record for record in records}
    summary = {
        "data_status": "Fresh fits of the supplied XLSX profiles; archived TXT fits are not reproduced.",
        "model": "I(x) = A exp(-2 (x - c)^2 / w^2) + b; w > 0",
        "position_units": "Sample index; nominal pixels, with sampling and physical scale unverified.",
        "physical_radius_um": None,
        "uncertainty_note": "Residual-scaled covariance errors only; no calibration, model or clipping uncertainty.",
        "config": config,
        "environment": {"python": platform.python_version(), "numpy": np.__version__,
                        "scipy": scipy.__version__, "matplotlib": matplotlib.__version__,
                        "openpyxl": openpyxl.__version__},
        "profiles": records,
        "expansion": {
            "status": "Diagnostic sample-width ratios only; optical magnification is not validated.",
            "assumption": "The two profiles would need equal spatial scale for an optical width ratio.",
            "free_baseline": radius_ratio(by_id["before"]["free_baseline"], by_id["after"]["free_baseline"]),
            "zero_baseline_sensitivity": radius_ratio(by_id["before"]["zero_baseline_sensitivity"], by_id["after"]["zero_baseline_sensitivity"]),
            "archived_notebook_ratio": {"value": 2.63, "stderr": 0.06, "status": "Cached historical output, different unavailable inputs."},
            "nominal_design_ratio": 100 / 25.4,
        },
    }
    output.mkdir(parents=True, exist_ok=True)
    (output / "summary.json").write_text(json.dumps(summary, indent=2, allow_nan=False) + "\n")
    with (output / "fit-summary.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["profile", "samples", "radius_samples", "radius_stderr_samples",
                         "diameter_samples", "offset_counts", "rmse_counts", "samples_at_ceiling", "full_radius_covered"])
        for record in records:
            fit = record["free_baseline"]
            writer.writerow([record["id"]] + [fit[key] for key in ["samples", "radius_samples",
                "radius_stderr_samples", "diameter_samples", "offset_counts", "rmse_counts",
                "samples_at_ceiling", "full_radius_covered"]])
    for entry, x, y, free, zero in datasets:
        fitted = prediction(x, free)
        np.savetxt(output / f"{entry['id']}-profile.csv", np.column_stack((x, y, fitted, y - fitted)),
                   delimiter=",", header="sample_index,intensity_counts,fitted_counts,residual_counts",
                   comments="", fmt="%.10g")
    if make_figures:
        save_figures(datasets, output)
    return summary


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--config", type=Path, default=PROJECT / "config/analysis.json")
    parser.add_argument("--output-dir", type=Path, required=True,
                        help="Generated output folder; repeated runs replace generated files here.")
    args = parser.parse_args()
    try:
        result = run_analysis(args.config, args.output_dir)
    except (ValueError, OSError, KeyError, RuntimeError) as error:
        parser.exit(2, f"Analysis failed: {error}\n")
    for record in result["profiles"]:
        fit = record["free_baseline"]
        print(f"{record['id']}: w = {fit['radius_samples']:.4f} samples; RMSE = {fit['rmse_counts']:.3f} counts")
        for warning in fit["warnings"]:
            print(f"  {warning}")
    print(result["expansion"]["status"])
    print(f"Results written to {args.output_dir}")


if __name__ == "__main__":
    main()
